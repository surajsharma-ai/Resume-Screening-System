"""
AI Resume Screening System - Complete Flask Application
"""

from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify, send_file
from werkzeug.utils import secure_filename
import sqlite3
import json
import re
import os
import uuid
from datetime import datetime, timedelta
from io import BytesIO
from fpdf import FPDF
import bcrypt
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

# ==================== CONFIGURATION ====================
app = Flask(__name__)
app.secret_key = 'your-secret-key-12345'

DATABASE = 'data/resume.db'
UPLOAD_FOLDER = 'uploads'
ALLOWED_EXTENSIONS = {'pdf'}

os.makedirs('data', exist_ok=True)
os.makedirs('uploads', exist_ok=True)

# Skills database
SKILLS = [
    'python', 'java', 'javascript', 'c++', 'c#', 'html', 'css', 'react', 'angular', 
    'node.js', 'django', 'flask', 'sql', 'mysql', 'mongodb', 'aws', 'docker', 
    'kubernetes', 'machine learning', 'deep learning', 'nlp', 'tensorflow', 
    'pytorch', 'pandas', 'numpy', 'git', 'linux', 'api', 'rest', 'agile'
]

GENDER_WORDS = ['he', 'she', 'him', 'her', 'his', 'hers', 'male', 'female', 'mr', 'mrs', 'ms']

# ==================== DATABASE ====================
def get_db():
    conn = sqlite3.connect(DATABASE)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db()
    conn.executescript('''
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE,
            password TEXT,
            email TEXT,
            fullname TEXT,
            phone TEXT,
            role TEXT,
            company TEXT
        );
        CREATE TABLE IF NOT EXISTS jobs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            recruiter_id INTEGER,
            title TEXT,
            company TEXT,
            location TEXT,
            description TEXT,
            requirements TEXT,
            skills TEXT,
            experience INTEGER,
            salary TEXT,
            status TEXT DEFAULT 'active'
        );
        CREATE TABLE IF NOT EXISTS applications (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            job_id INTEGER,
            applicant_id INTEGER,
            resume_text TEXT,
            anon_text TEXT,
            match_score REAL,
            skill_score REAL,
            matched_skills TEXT,
            missing_skills TEXT,
            experience INTEGER,
            status TEXT DEFAULT 'pending',
            UNIQUE(job_id, applicant_id)
        );
        CREATE TABLE IF NOT EXISTS notifications (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER,
            title TEXT,
            message TEXT,
            type TEXT,
            is_read INTEGER DEFAULT 0
        );
        CREATE TABLE IF NOT EXISTS interview_questions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            job_id INTEGER,
            question TEXT,
            question_order INTEGER
        );
        CREATE TABLE IF NOT EXISTS interviews (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            application_id INTEGER,
            job_id INTEGER,
            applicant_id INTEGER,
            recruiter_id INTEGER,
            scheduled_date TEXT,
            scheduled_time TEXT,
            duration_minutes INTEGER DEFAULT 30,
            room_id TEXT UNIQUE,
            status TEXT DEFAULT 'scheduled',
            notes TEXT,
            rating INTEGER
        );
        CREATE TABLE IF NOT EXISTS interview_responses (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            application_id INTEGER,
            question_id INTEGER,
            applicant_id INTEGER,
            response TEXT,
            UNIQUE(application_id, question_id)
        );
    ''')
    conn.commit()
    conn.close()

init_db()

# ==================== HELPERS ====================
def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def hash_password(password):
    return bcrypt.hashpw(password.encode(), bcrypt.gensalt()).decode()

def check_password(password, hashed):
    return bcrypt.checkpw(password.encode(), hashed.encode())

def extract_text_from_file(filepath):
    """Extract text from uploaded resume"""
    ext = filepath.split('.')[-1].lower()
    try:
        if ext == 'txt':
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        elif ext == 'pdf':
            import PyPDF2
            with open(filepath, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                text = ''
                for page in reader.pages:
                    text += page.extract_text() or ''
                return text
        elif ext == 'docx':
            from docx import Document
            doc = Document(filepath)
            return ' '.join([p.text for p in doc.paragraphs])
    except:
        pass
    return ''

def anonymize_text(text):
    """Remove name, gender, contact info using REGEX"""
    # Remove emails
    text = re.sub(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', '[EMAIL]', text)
    # Remove phones
    text = re.sub(r'\+?\d{1,3}[-.\s]?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}', '[PHONE]', text)
    text = re.sub(r'\b\d{10,12}\b', '[PHONE]', text)
    # Remove gender words
    for word in GENDER_WORDS:
        text = re.sub(r'\b' + word + r'\b', '', text, flags=re.IGNORECASE)
    # Remove names (first 3 lines, capitalized words)
    lines = text.split('\n')
    skip = ['Resume', 'CV', 'Summary', 'Profile', 'Experience', 'Education', 'Skills']
    for i in range(min(3, len(lines))):
        for word in re.findall(r'\b[A-Z][a-z]{2,}\b', lines[i]):
            if word not in skip:
                lines[i] = lines[i].replace(word, '[NAME]')
    return '\n'.join(lines)

def calculate_match(resume_text, job_desc, required_skills):
    """Calculate match score using TF-IDF and skill matching"""
    # TF-IDF Similarity
    try:
        vectorizer = TfidfVectorizer(stop_words='english')
        matrix = vectorizer.fit_transform([resume_text.lower(), job_desc.lower()])
        similarity = cosine_similarity(matrix[0:1], matrix[1:2])[0][0]
    except:
        similarity = 0
    
    # Skill matching
    resume_lower = resume_text.lower()
    found_skills = [s for s in SKILLS if re.search(r'\b' + re.escape(s) + r'\b', resume_lower)]
    required_set = set(s.lower() for s in required_skills)
    matched = list(required_set & set(found_skills))
    missing = list(required_set - set(found_skills))
    skill_pct = (len(matched) / len(required_set) * 100) if required_set else 0
    
    # Experience
    exp_match = re.search(r'(\d+)\+?\s*(?:years?|yrs?)', resume_lower)
    experience = int(exp_match.group(1)) if exp_match else 0
    
    # Overall score
    overall = (similarity * 60) + (skill_pct * 0.4)
    
    return {
        'match_score': round(float(overall), 1),
        'skill_score': round(float(skill_pct), 1),
        'matched': matched,
        'missing': missing,
        'experience': experience,
        'all_skills': found_skills
    }

# ==================== ROUTES ====================

@app.route('/')
def landing():
    if 'user_id' in session:
        if session['role'] == 'recruiter':
            return redirect(url_for('recruiter_dashboard'))
        return redirect(url_for('applicant_dashboard'))
    return render_template('landing.html')

# -------- RECRUITER AUTH --------
@app.route('/recruiter/login', methods=['GET', 'POST'])
def recruiter_login():
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']
        
        conn = get_db()
        user = conn.execute('SELECT * FROM users WHERE username=? AND role=?', 
                           (username, 'recruiter')).fetchone()
        conn.close()
        
        if user and check_password(password, user['password']):
            session['user_id'] = user['id']
            session['username'] = user['username']
            session['fullname'] = user['fullname']
            session['role'] = 'recruiter'
            session['company'] = user['company']
            flash('Login successful!', 'success')
            return redirect(url_for('recruiter_dashboard'))
        flash('Invalid credentials', 'danger')
    return render_template('recruiter_login.html')

@app.route('/recruiter/register', methods=['GET', 'POST'])
def recruiter_register():
    if request.method == 'POST':
        try:
            conn = get_db()
            conn.execute('''INSERT INTO users (username, password, email, fullname, phone, role, company)
                           VALUES (?, ?, ?, ?, ?, ?, ?)''',
                        (request.form['username'], hash_password(request.form['password']),
                         request.form['email'], request.form['fullname'], 
                         request.form['phone'], 'recruiter', request.form['company']))
            conn.commit()
            conn.close()
            flash('Registration successful! Please login.', 'success')
            return redirect(url_for('recruiter_login'))
        except sqlite3.IntegrityError:
            flash('Username already exists', 'danger')
    return render_template('recruiter_register.html')

# -------- APPLICANT AUTH --------
@app.route('/applicant/login', methods=['GET', 'POST'])
def applicant_login():
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']
        
        conn = get_db()
        user = conn.execute('SELECT * FROM users WHERE username=? AND role=?', 
                           (username, 'applicant')).fetchone()
        conn.close()
        
        if user and check_password(password, user['password']):
            session['user_id'] = user['id']
            session['username'] = user['username']
            session['fullname'] = user['fullname']
            session['role'] = 'applicant'
            flash('Login successful!', 'success')
            return redirect(url_for('applicant_dashboard'))
        flash('Invalid credentials', 'danger')
    return render_template('applicant_login.html')

@app.route('/applicant/register', methods=['GET', 'POST'])
def applicant_register():
    if request.method == 'POST':
        try:
            conn = get_db()
            conn.execute('''INSERT INTO users (username, password, email, fullname, phone, role)
                           VALUES (?, ?, ?, ?, ?, ?)''',
                        (request.form['username'], hash_password(request.form['password']),
                         request.form['email'], request.form['fullname'], 
                         request.form['phone'], 'applicant'))
            conn.commit()
            conn.close()
            flash('Registration successful! Please login.', 'success')
            return redirect(url_for('applicant_login'))
        except sqlite3.IntegrityError:
            flash('Username already exists', 'danger')
    return render_template('applicant_register.html')

@app.route('/logout')
def logout():
    session.clear()
    flash('Logged out', 'success')
    return redirect(url_for('landing'))

# -------- RECRUITER PAGES --------
@app.route('/recruiter/dashboard')
def recruiter_dashboard():
    if session.get('role') != 'recruiter':
        return redirect(url_for('landing'))
    
    conn = get_db()
    jobs = conn.execute('''SELECT j.*, 
                          (SELECT COUNT(*) FROM applications WHERE job_id=j.id) as app_count
                          FROM jobs j WHERE recruiter_id=? ORDER BY id DESC''', 
                       (session['user_id'],)).fetchall()
    conn.close()
    
    total_apps = sum(j['app_count'] for j in jobs)
    return render_template('recruiter_dashboard.html', jobs=jobs, total_apps=total_apps)

@app.route('/recruiter/post-job', methods=['GET', 'POST'])
def post_job():
    if session.get('role') != 'recruiter':
        return redirect(url_for('landing'))
    
    if request.method == 'POST':
        skills = [s.strip().lower() for s in request.form['skills'].split(',') if s.strip()]
        
        conn = get_db()
        conn.execute('''INSERT INTO jobs (recruiter_id, title, company, location, description,
                       requirements, skills, experience, salary)
                       VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                    (session['user_id'], request.form['title'], request.form['company'],
                     request.form['location'], request.form['description'],
                     request.form['requirements'], json.dumps(skills),
                     int(request.form.get('experience', 0)), request.form.get('salary', '')))
        conn.commit()
        conn.close()
        flash('Job posted successfully!', 'success')
        return redirect(url_for('recruiter_dashboard'))
    
    return render_template('post_job.html', company=session.get('company', ''))

@app.route('/recruiter/delete-job/<int:job_id>', methods=['POST'])
def delete_job(job_id):
    if session.get('role') != 'recruiter':
        return jsonify({'success': False})
    
    conn = get_db()
    job = conn.execute('SELECT * FROM jobs WHERE id=? AND recruiter_id=?', 
                       (job_id, session['user_id'])).fetchone()
    if job:
        conn.execute('DELETE FROM applications WHERE job_id=?', (job_id,))
        conn.execute('DELETE FROM jobs WHERE id=?', (job_id,))
        conn.commit()
        conn.close()
        return jsonify({'success': True, 'message': 'Job and all related applications deleted successfully!'})
    conn.close()
    return jsonify({'success': False, 'message': 'Job not found or access denied'})

@app.route('/recruiter/applications/<int:job_id>')
def view_applications(job_id):
    if session.get('role') != 'recruiter':
        return redirect(url_for('landing'))
    
    conn = get_db()
    job = conn.execute('SELECT * FROM jobs WHERE id=?', (job_id,)).fetchone()
    apps = conn.execute('''SELECT * FROM applications WHERE job_id=? 
                          ORDER BY match_score DESC''', (job_id,)).fetchall()
    conn.close()
    
    if not job:
        flash('Job not found', 'danger')
        return redirect(url_for('recruiter_dashboard'))
    
    job = dict(job)
    job['skills'] = json.loads(job['skills'])
    
    applications = []
    for i, a in enumerate(apps):
        app = dict(a)
        app['rank'] = i + 1
        app['matched_skills'] = json.loads(str(app['matched_skills']) if app['matched_skills'] else '[]')
        app['missing_skills'] = json.loads(str(app['missing_skills']) if app['missing_skills'] else '[]')
        applications.append(app)
    
    return render_template('view_applications.html', job=job, applications=applications)

@app.route('/recruiter/details/<int:app_id>')
def get_details(app_id):
    if session.get('role') != 'recruiter':
        return jsonify({'success': False})
    
    conn = get_db()
    result = conn.execute('''SELECT u.fullname, u.email, u.phone, a.resume_text
                            FROM applications a JOIN users u ON a.applicant_id=u.id
                            WHERE a.id=?''', (app_id,)).fetchone()
    conn.close()
    
    if result:
        return jsonify({
            'success': True,
            'name': result['fullname'],
            'email': result['email'],
            'phone': result['phone'] or 'Not provided',
            'resume': result['resume_text'][:2000] if result['resume_text'] else '',
            'app_id': app_id
        })
    return jsonify({'success': False})

@app.route('/recruiter/download-resume/<int:app_id>')
def download_resume(app_id):
    if session.get('role') != 'recruiter':
        return redirect(url_for('landing'))
    
    conn = get_db()
    result = conn.execute('''SELECT u.fullname, a.resume_text, a.applicant_id, a.job_id
                            FROM applications a JOIN users u ON a.applicant_id=u.id
                            WHERE a.id=?''', (app_id,)).fetchone()
    conn.close()
    
    if not result:
        flash('Application not found', 'danger')
        return redirect(url_for('recruiter_dashboard'))

    # Check multiple naming patterns for the original PDF
    possible_filenames = [
        f"resume_{result['applicant_id']}_{result['job_id']}.pdf",
        f"{result['applicant_id']}_{result['job_id']}_*.pdf" # Glob-style
    ]
    
    filepath = None
    import glob
    for pattern in possible_filenames:
        matches = glob.glob(os.path.join(UPLOAD_FOLDER, pattern))
        if matches:
            filepath = matches[0]
            break

    safe_name = re.sub(r'[^a-zA-Z0-9_]', '_', result['fullname'])

    # Serve original PDF if it exists
    if filepath and os.path.exists(filepath):
        return send_file(filepath, as_attachment=True, 
                         download_name=f'Resume_{safe_name}.pdf',
                         mimetype='application/pdf')
    
    # Fallback to text-to-pdf if original file is missing
    if not result['resume_text']:
        flash('Resume content not found', 'danger')
        return redirect(url_for('recruiter_dashboard'))
    
    import textwrap
    
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font('Helvetica', 'B', 16)
    # Using a fixed width (190mm for A4) instead of 0 to avoid "horizontal space" errors
    pdf.cell(190, 10, f"Resume - {result['fullname']}", new_x='LMARGIN', new_y='NEXT', align='C')
    pdf.ln(5)
    pdf.set_font('Helvetica', '', 11)
    
    for line in result['resume_text'].split('\n'):
        clean_line = line.encode('latin-1', 'replace').decode('latin-1')
        wrapped = textwrap.wrap(clean_line, width=70, break_long_words=True, replace_whitespace=False)
        if not wrapped:
            pdf.ln(6)
        else:
            for w in wrapped:
                # Use fixed width 190 instead of 0 for safety
                pdf.multi_cell(190, 6, w)
    
    buffer = BytesIO()
    pdf.output(buffer)
    buffer.seek(0)
    
    return send_file(buffer, as_attachment=True, 
                     download_name=f'Resume_{safe_name}.pdf',
                     mimetype='application/pdf')

@app.route('/recruiter/select/<int:app_id>', methods=['POST'])
def select_next_round(app_id):
    if session.get('role') != 'recruiter':
        return jsonify({'success': False})
    
    conn = get_db()
    application = conn.execute('SELECT * FROM applications WHERE id=?', (app_id,)).fetchone()
    if application:
        conn.execute('UPDATE applications SET status=? WHERE id=?', ('selected', app_id))
        job = conn.execute('SELECT title, company FROM jobs WHERE id=?', (application['job_id'],)).fetchone()
        conn.execute('''INSERT INTO notifications (user_id, title, message, type)
                       VALUES (?, ?, ?, ?)''',
                    (application['applicant_id'], '🎉 Selected for Next Round!',
                     f'Congratulations! You have been selected for the next round for {job["title"]} at {job["company"]}. Please check your applications dashboard to answer the screening questions.', 'success'))
        conn.commit()
    conn.close()
    return jsonify({'success': True, 'message': 'Candidate selected for next round! Notification sent.'})

@app.route('/recruiter/reject/<int:app_id>', methods=['POST'])
def reject(app_id):
    if session.get('role') != 'recruiter':
        return jsonify({'success': False})
    
    conn = get_db()
    application = conn.execute('SELECT * FROM applications WHERE id=?', (app_id,)).fetchone()
    if application:
        conn.execute('UPDATE applications SET status=? WHERE id=?', ('rejected', app_id))
        job = conn.execute('SELECT title, company FROM jobs WHERE id=?', (application['job_id'],)).fetchone()
        conn.execute('''INSERT INTO notifications (user_id, title, message, type)
                       VALUES (?, ?, ?, ?)''',
                    (application['applicant_id'], 'Application Update',
                     f'Thank you for applying to {job["title"]}. We have decided to proceed with other candidates.', 'danger'))
        conn.commit()
    conn.close()
    return jsonify({'success': True, 'message': 'Candidate rejected. Notification sent.'})

@app.route('/recruiter/job/<int:job_id>/questions', methods=['GET', 'POST'])
def manage_questions(job_id):
    if session.get('role') != 'recruiter':
        return redirect(url_for('landing'))
    
    conn = get_db()
    job = conn.execute('SELECT * FROM jobs WHERE id=? AND recruiter_id=?',
                       (job_id, session['user_id'])).fetchone()
    if not job:
        flash('Job not found', 'danger')
        return redirect(url_for('recruiter_dashboard'))
    
    if request.method == 'POST':
        action = request.form.get('action')
        if action == 'add':
            question = request.form.get('question', '').strip()
            if question:
                max_order = conn.execute('SELECT MAX(question_order) FROM interview_questions WHERE job_id=?',
                                         (job_id,)).fetchone()[0]
                next_order = (max_order or 0) + 1
                conn.execute('INSERT INTO interview_questions (job_id, question, question_order) VALUES (?, ?, ?)',
                            (job_id, question, next_order))
                conn.commit()
                flash('Question added!', 'success')
        elif action == 'delete':
            q_id = request.form.get('question_id')
            conn.execute('DELETE FROM interview_questions WHERE id=? AND job_id=?', (q_id, job_id))
            conn.commit()
            flash('Question removed.', 'success')
    
    questions = conn.execute('SELECT * FROM interview_questions WHERE job_id=? ORDER BY question_order',
                             (job_id,)).fetchall()
    conn.close()
    return render_template('interview_questions.html', job=job, questions=questions)

@app.route('/recruiter/schedule-interview/<int:app_id>', methods=['POST'])
def schedule_interview(app_id):
    if session.get('role') != 'recruiter':
        return jsonify({'success': False})
    
    conn = get_db()
    application = conn.execute('SELECT * FROM applications WHERE id=?', (app_id,)).fetchone()
    if not application:
        conn.close()
        return jsonify({'success': False, 'message': 'Application not found'})
    
    scheduled_date = request.form.get('date')
    scheduled_time = request.form.get('time')
    duration = int(request.form.get('duration', 30))
    room_id = str(uuid.uuid4())[:8]
    
    conn.execute('''INSERT INTO interviews 
                    (application_id, job_id, applicant_id, recruiter_id, scheduled_date, scheduled_time, duration_minutes, room_id)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)''',
                (app_id, application['job_id'], application['applicant_id'],
                 session['user_id'], scheduled_date, scheduled_time, duration, room_id))
    conn.execute('UPDATE applications SET status=? WHERE id=?', ('interview_scheduled', app_id))
    
    job = conn.execute('SELECT title, company FROM jobs WHERE id=?', (application['job_id'],)).fetchone()
    conn.execute('''INSERT INTO notifications (user_id, title, message, type)
                   VALUES (?, ?, ?, ?)''',
                (application['applicant_id'], '📅 Interview Scheduled!',
                 f'Your interview for {job["title"]} at {job["company"]} is scheduled on {scheduled_date} at {scheduled_time}. Check your applications to join the interview room.', 'info'))
    conn.commit()
    conn.close()
    return jsonify({'success': True, 'message': f'Interview scheduled for {scheduled_date} at {scheduled_time}!'})

@app.route('/interview-room/<room_id>')
def interview_room(room_id):
    if 'user_id' not in session:
        return redirect(url_for('landing'))
    
    conn = get_db()
    interview = conn.execute('SELECT * FROM interviews WHERE room_id=?', (room_id,)).fetchone()
    if not interview:
        flash('Interview room not found', 'danger')
        conn.close()
        return redirect(url_for('landing'))
    
    # Check if interview has started
    scheduled_dt_str = f"{interview['scheduled_date']} {interview['scheduled_time']}"
    try:
        scheduled_dt = datetime.strptime(scheduled_dt_str, '%Y-%m-%d %H:%M')
        if datetime.now() < scheduled_dt:
            conn.close()
            return render_template('interview_waiting.html', 
                                   date=interview['scheduled_date'], 
                                   time=interview['scheduled_time'],
                                   role=session.get('role'))
    except ValueError:
        pass # Fallback if time format is unexpected
    
    # Verify user is either the recruiter or the applicant
    user_id = session['user_id']
    role = session['role']
    if role == 'recruiter' and interview['recruiter_id'] != user_id:
        flash('Access denied', 'danger')
        conn.close()
        return redirect(url_for('recruiter_dashboard'))
    if role == 'applicant' and interview['applicant_id'] != user_id:
        flash('Access denied', 'danger')
        conn.close()
        return redirect(url_for('applicant_dashboard'))
    
    job = conn.execute('SELECT * FROM jobs WHERE id=?', (interview['job_id'],)).fetchone()
    applicant = conn.execute('SELECT fullname, email FROM users WHERE id=?',
                             (interview['applicant_id'],)).fetchone()
    recruiter = conn.execute('SELECT fullname, company FROM users WHERE id=?',
                             (interview['recruiter_id'],)).fetchone()
    questions = conn.execute('SELECT * FROM interview_questions WHERE job_id=? ORDER BY question_order',
                             (interview['job_id'],)).fetchall()
    conn.close()
    
    return render_template('interview_room.html',
                           interview=interview, job=job,
                           applicant=applicant, recruiter=recruiter,
                           questions=questions, role=role)

@app.route('/recruiter/interview-feedback/<int:interview_id>', methods=['POST'])
def interview_feedback(interview_id):
    if session.get('role') != 'recruiter':
        return jsonify({'success': False})
    
    conn = get_db()
    interview = conn.execute('SELECT * FROM interviews WHERE id=? AND recruiter_id=?',
                             (interview_id, session['user_id'])).fetchone()
    if interview:
        notes = request.form.get('notes', '')
        rating = int(request.form.get('rating', 0))
        conn.execute('UPDATE interviews SET notes=?, rating=?, status=? WHERE id=?',
                    (notes, rating, 'completed', interview_id))
        conn.execute('UPDATE applications SET status=? WHERE id=?',
                    ('interviewed', interview['application_id']))
        
        job = conn.execute('SELECT title, company FROM jobs WHERE id=?', (interview['job_id'],)).fetchone()
        conn.execute('''INSERT INTO notifications (user_id, title, message, type)
                       VALUES (?, ?, ?, ?)''',
                    (interview['applicant_id'], '🎤 Interview Completed',
                     f'Your interview for {job["title"]} at {job["company"]} has been completed. Results will be shared soon.', 'info'))
        conn.commit()
    conn.close()
    return jsonify({'success': True, 'message': 'Feedback saved successfully!'})

@app.route('/recruiter/final-hire/<int:app_id>', methods=['POST'])
def final_hire(app_id):
    if session.get('role') != 'recruiter':
        return jsonify({'success': False})
    
    conn = get_db()
    application = conn.execute('SELECT * FROM applications WHERE id=?', (app_id,)).fetchone()
    if application:
        conn.execute('UPDATE applications SET status=? WHERE id=?', ('hired', app_id))
        job = conn.execute('SELECT title, company FROM jobs WHERE id=?', (application['job_id'],)).fetchone()
        conn.execute('''INSERT INTO notifications (user_id, title, message, type)
                       VALUES (?, ?, ?, ?)''',
                    (application['applicant_id'], '🎉 Congratulations! You are HIRED!',
                     f'You have been officially hired for {job["title"]} at {job["company"]}! Welcome aboard!', 'success'))
        conn.commit()
    conn.close()
    return jsonify({'success': True, 'message': 'Candidate hired! Notification sent.'})

@app.route('/recruiter/get-interview/<int:app_id>')
def get_interview_info(app_id):
    if session.get('role') != 'recruiter':
        return jsonify({'success': False})
    
    conn = get_db()
    interview = conn.execute('SELECT * FROM interviews WHERE application_id=? ORDER BY id DESC LIMIT 1',
                             (app_id,)).fetchone()
    conn.close()
    
    if interview:
        return jsonify({
            'success': True,
            'interview_id': interview['id'],
            'room_id': interview['room_id'],
            'date': interview['scheduled_date'],
            'time': interview['scheduled_time'],
            'duration': interview['duration_minutes'],
            'status': interview['status'],
            'notes': interview['notes'] or '',
            'rating': interview['rating'] or 0
        })
    return jsonify({'success': False})

@app.route('/recruiter/view-responses/<int:app_id>')
def view_responses(app_id):
    if session.get('role') != 'recruiter':
        return jsonify({'success': False})
    
    conn = get_db()
    responses = conn.execute('''SELECT ir.response, iq.question, iq.question_order
                               FROM interview_responses ir
                               JOIN interview_questions iq ON ir.question_id=iq.id
                               WHERE ir.application_id=?
                               ORDER BY iq.question_order''', (app_id,)).fetchall()
    conn.close()
    
    if responses:
        data = [{'question': r['question'], 'response': r['response'], 'order': r['question_order']} for r in responses]
        return jsonify({'success': True, 'responses': data})
    return jsonify({'success': False, 'message': 'No responses found'})

# -------- APPLICANT PAGES --------
@app.route('/applicant/dashboard')
def applicant_dashboard():
    if session.get('role') != 'applicant':
        return redirect(url_for('landing'))
    
    conn = get_db()
    notif_count = conn.execute('SELECT COUNT(*) FROM notifications WHERE user_id=? AND is_read=0',
                               (session['user_id'],)).fetchone()[0]
    conn.close()
    return render_template('applicant_dashboard.html', notif_count=notif_count)

@app.route('/applicant/jobs')
def browse_jobs():
    if session.get('role') != 'applicant':
        return redirect(url_for('landing'))
    
    conn = get_db()
    jobs = conn.execute('SELECT * FROM jobs WHERE status=? ORDER BY id DESC', ('active',)).fetchall()
    applied = conn.execute('SELECT job_id FROM applications WHERE applicant_id=?',
                          (session['user_id'],)).fetchall()
    conn.close()
    
    applied_ids = set(a['job_id'] for a in applied)
    
    job_list = []
    for j in jobs:
        job = dict(j)
        job['skills'] = json.loads(str(job['skills']) if job['skills'] else '[]')
        job_list.append(job)
    
    return render_template('browse_jobs.html', jobs=job_list, applied_ids=applied_ids)

@app.route('/applicant/apply/<int:job_id>', methods=['GET', 'POST'])
def apply_job(job_id):
    if session.get('role') != 'applicant':
        return redirect(url_for('landing'))
    
    conn = get_db()
    job = conn.execute('SELECT * FROM jobs WHERE id=?', (job_id,)).fetchone()
    conn.close()
    
    if not job:
        flash('Job not found', 'danger')
        return redirect(url_for('browse_jobs'))
    
    job = dict(job)
    job['skills'] = json.loads(str(job['skills']) if job['skills'] else '[]')
    
    result = None
    applied = False
    
    if request.method == 'POST':
        if 'resume' not in request.files:
            flash('No file uploaded', 'danger')
        else:
            file = request.files['resume']
            if file and allowed_file(file.filename):
                # Save as a deterministic filename to easily retrieve later
                filename = f"resume_{session['user_id']}_{job_id}.pdf"
                filepath = os.path.join(UPLOAD_FOLDER, filename)
                file.save(filepath)
                
                # Extract text
                resume_text = extract_text_from_file(filepath)
                if not resume_text:
                    flash('Could not read resume', 'danger')
                else:
                    # Anonymize
                    anon_text = anonymize_text(resume_text)
                    
                    # Calculate match
                    jd = f"{job['description']} {job['requirements']}"
                    result = calculate_match(anon_text, jd, job['skills'])
                    
                    # Save application
                    try:
                        conn = get_db()
                        conn.execute('''INSERT INTO applications 
                                       (job_id, applicant_id, resume_text, anon_text, match_score,
                                        skill_score, matched_skills, missing_skills, experience)
                                       VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                                    (job_id, session['user_id'], resume_text, anon_text,
                                     result['match_score'], result['skill_score'],
                                     json.dumps(result['matched']), json.dumps(result['missing']),
                                     result['experience']))
                        conn.commit()
                        conn.close()
                        flash('Application submitted!', 'success')
                        applied = True
                    except sqlite3.IntegrityError:
                        flash('Already applied to this job', 'warning')
            else:
                flash('Invalid file type', 'danger')
    
    return render_template('apply_job.html', job=job, result=result, applied=applied)

@app.route('/applicant/my-applications')
def my_applications():
    if session.get('role') != 'applicant':
        return redirect(url_for('landing'))
    
    conn = get_db()
    # INNER JOIN ensures orphaned applications (where job no longer exists) are not shown
    # LEFT JOIN interviews to get interview details
    apps = conn.execute('''SELECT a.*, j.title, j.company,
                          i.scheduled_date as interview_date, i.scheduled_time as interview_time,
                          i.room_id as interview_room_id
                          FROM applications a
                          INNER JOIN jobs j ON a.job_id=j.id
                          LEFT JOIN interviews i ON i.application_id=a.id
                          WHERE a.applicant_id=? ORDER BY a.id DESC''',
                       (session['user_id'],)).fetchall()
    conn.close()
    
    applications = []
    for a in apps:
        app = dict(a)
        app['matched_skills'] = json.loads(str(app['matched_skills']) if app['matched_skills'] else '[]')
        app['missing_skills'] = json.loads(str(app['missing_skills']) if app['missing_skills'] else '[]')
        applications.append(app)
    
    return render_template('my_applications.html', applications=applications)

@app.route('/applicant/delete-application/<int:app_id>', methods=['POST'])
def delete_application(app_id):
    if session.get('role') != 'applicant':
        return jsonify({'success': False})
    
    conn = get_db()
    application = conn.execute('SELECT * FROM applications WHERE id=? AND applicant_id=?',
                               (app_id, session['user_id'])).fetchone()
    if application:
        conn.execute('DELETE FROM applications WHERE id=?', (app_id,))
        conn.commit()
        conn.close()
        return jsonify({'success': True, 'message': 'Application deleted successfully!'})
    conn.close()
    return jsonify({'success': False, 'message': 'Application not found or access denied'})

@app.route('/applicant/answer-questions/<int:app_id>', methods=['GET', 'POST'])
def answer_questions(app_id):
    if session.get('role') != 'applicant':
        return redirect(url_for('landing'))
    
    conn = get_db()
    application = conn.execute('SELECT * FROM applications WHERE id=? AND applicant_id=?',
                               (app_id, session['user_id'])).fetchone()
    if not application:
        flash('Application not found', 'danger')
        conn.close()
        return redirect(url_for('my_applications'))
    
    job = conn.execute('SELECT * FROM jobs WHERE id=?', (application['job_id'],)).fetchone()
    questions = conn.execute('SELECT * FROM interview_questions WHERE job_id=? ORDER BY question_order',
                             (application['job_id'],)).fetchall()
    
    if not questions:
        flash('No screening questions for this position.', 'info')
        conn.close()
        return redirect(url_for('my_applications'))
    
    if request.method == 'POST':
        for q in questions:
            response = request.form.get(f'response_{q["id"]}', '').strip()
            if response:
                try:
                    conn.execute('''INSERT INTO interview_responses (application_id, question_id, applicant_id, response)
                                   VALUES (?, ?, ?, ?)''',
                                (app_id, q['id'], session['user_id'], response))
                except Exception:
                    conn.execute('''UPDATE interview_responses SET response=?
                                   WHERE application_id=? AND question_id=?''',
                                (response, app_id, q['id']))
        
        conn.execute('UPDATE applications SET status=? WHERE id=?', ('questions_answered', app_id))
        
        # Notify the recruiter
        applicant = conn.execute('SELECT fullname FROM users WHERE id=?', (session['user_id'],)).fetchone()
        conn.execute('''INSERT INTO notifications (user_id, title, message, type)
                       VALUES (?, ?, ?, ?)''',
                    (job['recruiter_id'], 'Screening Questions Answered',
                     f"{applicant['fullname']} has submitted answers to the screening questions for {job['title']}. You can now review them and schedule an interview.", 'info'))
                     
        conn.commit()
        flash('Responses saved! Waiting for recruiter to schedule the interview.', 'success')
        conn.close()
        return redirect(url_for('my_applications'))
    
    # Load existing responses if any
    existing = {}
    responses = conn.execute('SELECT * FROM interview_responses WHERE application_id=?', (app_id,)).fetchall()
    for r in responses:
        existing[r['question_id']] = r['response']
    conn.close()
    
    return render_template('answer_questions.html', job=job, questions=questions, 
                           app_id=app_id, existing=existing)

@app.route('/recruiter/reschedule-interview/<int:app_id>', methods=['POST'])
def reschedule_interview(app_id):
    if session.get('role') != 'recruiter':
        return jsonify({'success': False})
    
    conn = get_db()
    application = conn.execute('SELECT * FROM applications WHERE id=?',
                               (app_id,)).fetchone()
    if not application:
        conn.close()
        return jsonify({'success': False, 'message': 'Application not found'})
    
    interview = conn.execute('SELECT * FROM interviews WHERE application_id=? ORDER BY id DESC LIMIT 1',
                             (app_id,)).fetchone()
    if not interview:
        conn.close()
        return jsonify({'success': False, 'message': 'No interview found'})
    
    new_date = request.form.get('date')
    new_time = request.form.get('time')
    
    old_date = interview['scheduled_date']
    old_time = interview['scheduled_time']
    
    conn.execute('UPDATE interviews SET scheduled_date=?, scheduled_time=? WHERE id=?',
                (new_date, new_time, interview['id']))
    
    job = conn.execute('SELECT title, company FROM jobs WHERE id=?', (application['job_id'],)).fetchone()
    recruiter = conn.execute('SELECT fullname FROM users WHERE id=?', (session['user_id'],)).fetchone()
    
    # Notify applicant about reschedule
    conn.execute('''INSERT INTO notifications (user_id, title, message, type)
                   VALUES (?, ?, ?, ?)''',
                (application['applicant_id'], 'Interview Rescheduled',
                 f'Your interview for {job["title"]} at {job["company"]} has been rescheduled by {recruiter["fullname"]} from {old_date} at {old_time} to {new_date} at {new_time}.', 'warning'))
    
    # Notify recruiter confirmation
    conn.execute('''INSERT INTO notifications (user_id, title, message, type)
                   VALUES (?, ?, ?, ?)''',
                (session['user_id'], 'Interview Rescheduled Successfully',
                 f'You successfully rescheduled the interview to {new_date} at {new_time}.', 'info'))
    
    conn.commit()
    conn.close()
    return jsonify({'success': True, 'message': 'Interview rescheduled successfully!'})

@app.route('/applicant/notifications')
def notifications():
    if session.get('role') != 'applicant':
        return redirect(url_for('landing'))
    
    conn = get_db()
    notifs = conn.execute('SELECT * FROM notifications WHERE user_id=? ORDER BY id DESC',
                         (session['user_id'],)).fetchall()
    conn.execute('UPDATE notifications SET is_read=1 WHERE user_id=?', (session['user_id'],))
    conn.commit()
    conn.close()
    
    return render_template('notifications.html', notifications=notifs)

# ==================== RUN ====================
if __name__ == '__main__':
    print("\n" + "="*50)
    print("Starting AI Resume Screening System")
    print("="*50)
    print("\nOpen this URL in your browser:")
    print("   http://127.0.0.1:5000")
    print("\n" + "="*50 + "\n")
    app.run(debug=True, port=5000)